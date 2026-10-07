# -*- coding: utf-8 -*-
r"""
示例数据导入脚本（由「5-导入示例数据.bat」调用）

做四件事：
  1. 生成 5 份数码商品资料（覆盖 md / docx / xlsx / csv / txt 五种格式）
  2. 创建知识库「电商商品知识库」
  3. 走**真实入库流程**处理这 5 份资料（解析→分块→向量化→双库写入）
     —— 所以本脚本同时也是一次完整的入库功能测试
  4. 创建演示账号 demo / demo2（密码 demo12345）

注意：运行前请先关闭正在运行的系统（避免两个程序同时写向量库）。
重复运行不会重复导入（检测到已有示例知识库会提示跳过）。
"""
import os
import socket
import sys
from pathlib import Path

# 让脚本从任何位置运行都能 import 到 app 包
BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from app.config import DATA_DIR  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.database import Base, SessionLocal, engine  # noqa: E402
from app.models import Document, KnowledgeBase, User  # noqa: E402

DEMO_DIR = DATA_DIR / "demo_docs"
DEMO_KB_NAME = "电商商品知识库"
# 演示账号的公开密码（仅用于本地功能演示，非真实用户凭据）
DEMO_PASSWORD = os.environ.get("DEMO_PASSWORD") or "demo12345"


# ---------------------------------------------------------------------------
# 一、生成 5 份示例资料
# ---------------------------------------------------------------------------

_EARPHONE_MD = """# W1 真无线主动降噪耳机

## 基本信息
- 商品名称：W1 真无线主动降噪耳机
- 型号：W1
- 售价：399 元
- 可选颜色：曜石黑 / 云雾白
- 上市时间：2026 年 3 月

## 续航与充电
- 单次续航：开启降噪 22 小时，关闭降噪 30 小时（配合充电盒计算）
- 快充功能：充电 10 分钟，可连续听歌 2 小时
- 充电接口：USB-C
- 充电盒支持无线充电

## 音质与降噪
- 主动降噪深度：最高 45 分贝
- 支持通透模式（戴耳机也能听清环境声）
- 蓝牙 5.4，支持两台设备同时连接、无缝切换
- 支持音频编码：AAC / SBC / LDAC

## 防护与佩戴
- 防水等级：IPX5（可防汗、防小雨，不可游泳或浸泡水中）
- 单耳重量：4.8 克
- 佩戴方式：入耳式，随机附赠 S/M/L 三种尺寸耳塞

## 包装清单
耳机 ×2、充电盒 ×1、USB-C 充电线 ×1、替换耳塞 ×3 对、说明书 ×1

## 保修说明
整机保修 1 年，充电盒同样享受 1 年保修；支持 7 天无理由退货（未激活拆封为前提，详见售后政策文档）。
"""

_WATCH_TEXT = """S2 智能运动手表 商品资料

一、基本信息
商品名称：S2 智能运动手表
型号：S2
售价：899 元
屏幕：1.85 英寸 AMOLED 高清屏，支持息屏显示
机身重量：38 克（不含表带）
表带：亲肤硅胶，支持快拆更换

二、健康监测
支持血氧饱和度监测（SpO2）
支持全天候心率监测，静息心率异常自动提醒
支持睡眠监测与睡眠质量评分
支持压力监测与呼吸训练引导

三、运动功能
支持 120 种以上运动模式（跑步、游泳、骑行、瑜伽等）
内置双频 GPS 定位，轨迹记录更精准
防水等级：5ATM（可佩戴游泳，但不可用于热水淋浴、桑拿或潜水）

四、续航与充电
典型使用场景续航 14 天
重度使用（全天 GPS）续航约 8 天
磁吸式充电，约 90 分钟充满

五、兼容性
支持 Android 8.0 及以上系统
支持 iOS 12 及以上系统
蓝牙 5.2 连接
"""

_AFTER_SALE_TXT = """售后政策与常见问题（电商平台通用）

【退换货政策】
1. 7 天无理由退货：自签收之日起 7 天内，商品保持完好可申请无理由退货。
   注意：耳机、手表等个人贴身使用商品，一经激活或拆封使用，不支持无理由退货。
2. 15 天质量问题换货：签收后 15 天内出现非人为质量问题，可申请换新。
3. 退货时请保留全部包装、配件、赠品，缺件可能影响退款金额。

【保修政策】
1. 全品类商品整机保修 1 年（自签收之日起算）。
2. 人为损坏（进水、摔坏、私自拆机改装）不在保修范围内。
3. 保修期内维修免费，寄修运费需自行承担。
4. 过保商品可付费维修，具体报价以检测结果为准。

【运费规则】
1. 质量问题导致的退换货：运费由商家承担。
2. 7 天无理由退货：退回运费由顾客承担。
3. 单笔订单满 99 元包邮（偏远地区除外）。

【发票说明】
支持开具电子发票：下单时在"发票信息"处填写抬头与税号；
发票默认在确认收货后 3 个工作日内开出。

【其他常见问题】
Q：可以修改收货地址吗？
A：发货前可在订单详情页自行修改；已发货请联系客服协助。
Q：多久发货？
A：现货商品 48 小时内发出，预售商品以商品页面标注时间为准。
Q：支持哪些付款方式？
A：支持微信支付、支付宝、银行卡支付。
Q：售后申请入口在哪里？
A：我的订单 → 选择订单 → 申请售后，按页面提示操作即可。
"""

_ROBOT_CSV = """参数项,参数值
商品名称,R3 扫地机器人
型号,R3
售价,1299 元
吸力,8000Pa
电池容量,5200mAh
续航时间,约 150 分钟
工作噪音,65 分贝（安静模式）
导航方式,激光导航（LDS）
集尘盒容量,400 毫升
水箱容量,250 毫升
机身高度,9.7 厘米
自动回充,支持
断点续扫,支持
APP 控制,支持（手机远程控制与地图管理）
适用面积,120 平方米以内
"""


def _make_files() -> list[tuple[str, str]]:
    """生成 5 份示例文件，返回 [(文件名, 类型), ...]"""
    DEMO_DIR.mkdir(parents=True, exist_ok=True)
    created: list[tuple[str, str]] = []

    # 1. markdown
    md_name = "商品资料-无线耳机W1.md"
    (DEMO_DIR / md_name).write_text(_EARPHONE_MD, encoding="utf-8")
    created.append((md_name, "md"))

    # 2. word（段落 + 一张参数表格）
    docx_name = "商品资料-智能手表S2.docx"
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("S2 智能运动手表 商品资料", level=1)
    for line in _WATCH_TEXT.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "项目", "参数", "备注"
    for row_data in [
        ("屏幕", "1.85 英寸 AMOLED", "支持息屏显示"),
        ("续航", "14 天", "典型使用场景"),
        ("防水", "5ATM", "可游泳，忌热水"),
        ("重量", "38 克", "不含表带"),
    ]:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = val
    doc.save(DEMO_DIR / docx_name)
    created.append((docx_name, "docx"))

    # 3. excel（三款商品参数对比）
    xlsx_name = "商品参数对比表.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "商品参数对比"
    ws.append(["商品名称", "型号", "售价", "续航", "防水等级", "重量", "特色功能"])
    ws.append(
        ["W1 真无线降噪耳机", "W1", "399 元", "30 小时（关降噪）", "IPX5", "4.8 克（单耳）",
         "主动降噪 45dB、双设备连接"]
    )
    ws.append(
        ["S2 智能运动手表", "S2", "899 元", "14 天", "5ATM", "38 克",
         "血氧/心率监测、双频 GPS"]
    )
    ws.append(
        ["R3 扫地机器人", "R3", "1299 元", "150 分钟", "不支持", "3.2 千克",
         "8000Pa 吸力、激光导航、自动回充"]
    )
    wb.save(DEMO_DIR / xlsx_name)
    created.append((xlsx_name, "xlsx"))

    # 4. txt（售后政策）
    txt_name = "售后政策与常见问题.txt"
    (DEMO_DIR / txt_name).write_text(_AFTER_SALE_TXT, encoding="utf-8")
    created.append((txt_name, "txt"))

    # 5. csv（扫地机器人参数）
    csv_name = "扫地机器人R3参数.csv"
    (DEMO_DIR / csv_name).write_text(_ROBOT_CSV, encoding="utf-8")
    created.append((csv_name, "csv"))

    return created


# ---------------------------------------------------------------------------
# 二、导入流程
# ---------------------------------------------------------------------------


def _port_in_use(port: int = 8000) -> bool:
    """检测系统服务是否正在运行（避免两个程序同时写向量库）"""
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(1)
        return s.connect_ex(("127.0.0.1", port)) == 0


def main() -> int:
    print("=" * 50)
    print("  导入示例商品数据")
    print("=" * 50)

    if _port_in_use(8000):
        print()
        print("[提示] 检测到系统正在运行（端口 8000 被占用）。")
        print("       请先关闭「4-启动系统.bat」的窗口，再重新运行本脚本。")
        return 1

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        # 检查是否已导入过
        exist_kb = db.query(KnowledgeBase).filter(KnowledgeBase.name == DEMO_KB_NAME).first()
        if exist_kb is not None:
            print()
            print(f"[提示] 示例知识库「{DEMO_KB_NAME}」已存在，跳过导入。")
            print("       如果你想重新导入：先在管理界面删掉这个知识库，再运行本脚本。")
            _ensure_demo_users(db)
            return 0

        # 1. 生成文件
        print("\n[1/4] 生成示例商品资料（5 份，覆盖 5 种格式）...")
        files = _make_files()
        for name, ftype in files:
            print(f"      ✓ {name}")

        # 2. 建知识库
        print(f"\n[2/4] 创建知识库「{DEMO_KB_NAME}」...")
        kb = KnowledgeBase(
            name=DEMO_KB_NAME,
            description="示例数据：无线耳机 / 智能手表 / 扫地机器人 及售后政策",
        )
        db.add(kb)
        db.commit()
        db.refresh(kb)

        # 3. 建文档记录并逐份走真实入库流程
        print("\n[3/4] 处理文档（解析 → 分块 → 向量化入库）...")
        print("      第一次运行需要加载 AI 模型，可能等待 30 秒左右，请耐心")
        from app.services.ingestion_service import process_document  # 延迟导入

        ok_count = 0
        for name, ftype in files:
            file_path = DEMO_DIR / name
            doc = Document(
                kb_id=kb.id,
                filename=name,
                file_path=str(file_path),
                file_type=ftype,
                size_bytes=file_path.stat().st_size,
                status="pending",
            )
            db.add(doc)
            db.commit()
            db.refresh(doc)

            print(f"      → 正在处理：{name}")
            process_document(doc.id)  # 同步处理，能看到结果
            db.refresh(doc)
            if doc.status == "ready":
                ok_count += 1
                print(f"        完成：切分为 {doc.chunk_count} 段 ✓")
            else:
                print(f"        失败：{doc.error_msg}")

        # 4. 演示账号
        print("\n[4/4] 创建演示账号...")
        _ensure_demo_users(db)

        print()
        print("=" * 50)
        print(f"  导入完成！成功处理 {ok_count}/{len(files)} 份文档")
        print("  演示账号：demo / demo12345 、 demo2 / demo12345")
        print("  接下来双击「4-启动系统.bat」启动系统，")
        print("  用管理员登录进入「管理后台 → 知识库管理」查看结果")
        print("=" * 50)
        return 0 if ok_count == len(files) else 1
    finally:
        db.close()


def _ensure_demo_users(db) -> None:
    for name in ("demo", "demo2"):
        exist = db.query(User).filter(User.username == name).first()
        if exist is None:
            db.add(
                User(username=name, password_hash=hash_password(DEMO_PASSWORD), role="user")
            )
    db.commit()
    print(f"      演示账号就绪：demo / demo2（密码 {DEMO_PASSWORD}）")


if __name__ == "__main__":
    sys.exit(main())
